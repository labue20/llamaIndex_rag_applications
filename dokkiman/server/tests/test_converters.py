"""PDF to Word and Word to PDF conversion routes."""

import io
import subprocess
import tempfile
import zipfile
from pathlib import Path

import docx
import pytest
from werkzeug.datastructures import FileStorage

import pdf_to_word_service


# --- PDF to Word -------------------------------------------------------------

def test_pdf_to_word_converts_every_page(signup, make_pdf):
    pdf = make_pdf(pages=3, marker="Quarterly numbers")
    response = signup().post(
        "/convertPdfToWord",
        data={"file": (open(pdf, "rb"), "report.pdf")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "report_converted.docx" in response.headers["Content-Disposition"]

    text = "\n".join(p.text for p in docx.Document(io.BytesIO(response.data)).paragraphs)
    for page in (1, 2, 3):
        assert f"This is page {page}. Quarterly numbers for page {page}." in text


def test_pdf_to_word_rejects_non_pdf(signup):
    response = signup().post(
        "/convertPdfToWord",
        data={"file": (io.BytesIO(b"hello"), "notes.txt")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert response.get_json()["error"] == "File must be a PDF"


def test_pdf_to_word_rejects_corrupt_pdf(signup):
    response = signup().post(
        "/convertPdfToWord",
        data={"file": (io.BytesIO(b"not really a pdf"), "broken.pdf")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert response.get_json()["error"] == "Invalid PDF file or corrupted PDF"


def test_validate_pdf_file(monkeypatch):
    def upload(content, name="a.pdf"):
        return FileStorage(stream=io.BytesIO(content), filename=name)

    assert pdf_to_word_service.validate_pdf_file(upload(b"%PDF")) == (True, None)
    assert pdf_to_word_service.validate_pdf_file(upload(b"")) == (False, "Empty file provided")
    assert pdf_to_word_service.validate_pdf_file(upload(b"x", name="")) == (False, "No file selected")

    monkeypatch.setattr(pdf_to_word_service, "MAX_UPLOAD_BYTES", 10)
    monkeypatch.setattr(pdf_to_word_service, "MAX_UPLOAD_MB", 1)
    assert pdf_to_word_service.validate_pdf_file(upload(b"x" * 11)) == (
        False, "File is too large. The maximum size is 1 MB.")


# --- Word to PDF (route) ---------------------------------------------------
# Route tests replace the converter; LibreOffice itself is tested further down.

def _zip_bytes():
    """Minimal ZIP archive: passes the route's .docx check."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
    return buffer.getvalue()


FAKE_DOCX = _zip_bytes()


def _convert_word(user_client, filename, content=FAKE_DOCX):
    return user_client.post(
        "/convertWordToPdf",
        data={"file": (io.BytesIO(content), filename)},
        content_type="multipart/form-data",
    )


@pytest.fixture
def fake_converter(monkeypatch):
    """Replace the conversion and record the paths it was given."""
    import flask_demo

    seen = {}

    def fake_convert(input_path, output_path):
        seen["input"], seen["output"] = Path(input_path), Path(output_path)
        assert seen["input"].read_bytes() == FAKE_DOCX
        Path(output_path).write_bytes(b"%PDF-1.4 converted")
        return str(output_path)

    monkeypatch.setattr(flask_demo.converter, "convert_docx_to_pdf", fake_convert)
    return seen


def test_word_to_pdf_returns_the_converted_file(signup, fake_converter):
    response = _convert_word(signup(), "letter.docx")
    assert response.status_code == 200
    assert response.data == b"%PDF-1.4 converted"
    assert response.mimetype == "application/pdf"
    assert "letter_converted.pdf" in response.headers["Content-Disposition"]


def test_word_to_pdf_rejects_other_files(signup):
    user_client = signup()
    wrong_type = _convert_word(user_client, "slides.pptx")
    assert wrong_type.status_code == 400
    assert wrong_type.get_json()["error"] == "File must be a .docx document"
    assert user_client.post("/convertWordToPdf").status_code == 400


def test_word_to_pdf_rejects_files_that_are_not_really_docx(signup, fake_converter):
    response = _convert_word(signup(), "notes.docx", content=b"just some text, renamed")
    assert response.status_code == 400
    assert response.get_json()["error"] == "File must be a .docx document"
    assert fake_converter == {}  # never reached the converter


@pytest.mark.parametrize("filename", [
    "../../../etc/evil.docx",
    "..\\..\\windows\\evil.docx",
    "/absolute/path/report.docx",
    "../.docx",
])
def test_word_to_pdf_never_uses_the_raw_filename_in_paths(signup, fake_converter, filename):
    response = _convert_word(signup(), filename)

    assert response.status_code == 200
    download_name = response.headers["Content-Disposition"].split("filename=")[1]
    assert download_name.endswith("_converted.pdf")
    for unsafe in ("/", "\\", ".."):
        assert unsafe not in download_name
        assert unsafe not in fake_converter["input"].name
    # Both files lived directly inside one private temp folder
    work_dir = fake_converter["input"].parent
    assert work_dir.name.startswith("word-to-pdf-")
    assert work_dir.parent == Path(tempfile.gettempdir())
    assert fake_converter["output"].parent == work_dir


def test_word_to_pdf_keeps_readable_names(signup, fake_converter):
    response = _convert_word(signup(), "My Letter (final).docx")
    assert "My_Letter_final_converted.pdf" in response.headers["Content-Disposition"]


def test_word_to_pdf_deletes_its_files_after_success(signup, fake_converter):
    assert _convert_word(signup(), "letter.docx").status_code == 200
    assert not fake_converter["input"].parent.exists()


def test_word_to_pdf_deletes_its_files_after_failure(signup, monkeypatch):
    import flask_demo
    import word_to_pdf_service

    seen = {}

    def failing_convert(input_path, output_path):
        seen["input"] = Path(input_path)
        raise word_to_pdf_service.ConversionError("The document could not be converted.")

    monkeypatch.setattr(flask_demo.converter, "convert_docx_to_pdf", failing_convert)
    response = _convert_word(signup(), "letter.docx")

    assert response.status_code == 500
    assert "could not be converted" in response.get_json()["error"]
    assert not seen["input"].parent.exists()


@pytest.mark.parametrize("path", ["/convertWordToPdfInfo", "/downloadPdf/anything.pdf"])
def test_removed_word_to_pdf_routes_are_gone(signup, path):
    user_client = signup()
    assert user_client.post(path).status_code in (404, 405)
    assert user_client.get(path).status_code in (404, 405)


# --- Word to PDF (LibreOffice) -------------------------------------------------

@pytest.fixture
def docx_file(tmp_path):
    """A real two-page .docx file."""
    document = docx.Document()
    document.add_heading("Quarterly Report", 0)
    document.add_paragraph("First page: revenue was 4.2 million dollars.")
    document.add_page_break()
    document.add_paragraph("Second page: expenses were 1.1 million dollars.")
    path = tmp_path / "input" / "report.docx"
    path.parent.mkdir()
    document.save(path)
    return path


@pytest.fixture
def fake_soffice(monkeypatch, tmp_path):
    """Pretend LibreOffice is installed and record how it is run."""
    import word_to_pdf_service as svc

    soffice = tmp_path / "soffice"
    soffice.write_text("")
    monkeypatch.setattr(svc, "find_soffice", lambda: str(soffice))
    runs = []

    def run(command, **options):
        runs.append((command, options))
        out_dir = Path(command[command.index("--outdir") + 1])
        (out_dir / (Path(command[-1]).stem + ".pdf")).write_bytes(b"%PDF fake")
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(svc.subprocess, "run", run)
    return runs


def test_libreoffice_is_run_headless_with_its_own_profile(fake_soffice, docx_file, tmp_path):
    import word_to_pdf_service as svc

    output = tmp_path / "result.pdf"
    assert svc.converter.convert_docx_to_pdf(docx_file, output) == str(output)
    assert output.read_bytes() == b"%PDF fake"

    command, options = fake_soffice[0]
    assert "--headless" in command
    assert command[command.index("--convert-to") + 1] == "pdf"
    assert any(arg.startswith("-env:UserInstallation=file://") for arg in command)
    assert command[-1] == str(docx_file)
    assert options["timeout"] == svc.CONVERSION_TIMEOUT_SECONDS


def test_libreoffice_failure_is_reported(monkeypatch, fake_soffice, docx_file, tmp_path):
    import word_to_pdf_service as svc

    monkeypatch.setattr(svc.subprocess, "run",
                        lambda command, **o: subprocess.CompletedProcess(command, 1, "", "source file could not be loaded"))
    with pytest.raises(svc.ConversionError, match="could not be converted"):
        svc.converter.convert_docx_to_pdf(docx_file, tmp_path / "out.pdf")


def test_libreoffice_timeout_is_reported(monkeypatch, fake_soffice, docx_file, tmp_path):
    import word_to_pdf_service as svc

    def hang(command, **options):
        raise subprocess.TimeoutExpired(command, options["timeout"])

    monkeypatch.setattr(svc.subprocess, "run", hang)
    with pytest.raises(svc.ConversionError, match="longer than"):
        svc.converter.convert_docx_to_pdf(docx_file, tmp_path / "out.pdf")


def test_missing_libreoffice_is_reported(monkeypatch, docx_file, tmp_path):
    import word_to_pdf_service as svc

    monkeypatch.setattr(svc, "find_soffice", lambda: None)
    with pytest.raises(svc.ConversionError, match="LibreOffice is not installed"):
        svc.converter.convert_docx_to_pdf(docx_file, tmp_path / "out.pdf")


def test_find_soffice_prefers_the_configured_path(monkeypatch, tmp_path):
    import word_to_pdf_service as svc

    custom = tmp_path / "my-soffice"
    custom.write_text("")
    monkeypatch.setenv("SOFFICE_PATH", str(custom))
    assert svc.find_soffice() == str(custom)


@pytest.mark.skipif(
    __import__("word_to_pdf_service").find_soffice() is None,
    reason="LibreOffice is not installed",
)
def test_real_conversion_with_libreoffice(docx_file, tmp_path):
    import fitz
    import word_to_pdf_service as svc

    output = tmp_path / "report.pdf"
    svc.converter.convert_docx_to_pdf(docx_file, output)

    pdf = fitz.open(output)
    assert pdf.page_count == 2
    assert "revenue was 4.2 million dollars" in pdf[0].get_text()
    assert "expenses were 1.1 million dollars" in pdf[1].get_text()


def test_stale_libreoffice_profiles_are_removed(tmp_path, monkeypatch):
    import os
    import word_to_pdf_service as svc

    dead = tmp_path / "libreoffice-profile-999999"
    alive = tmp_path / f"libreoffice-profile-{os.getppid()}"
    unrelated = tmp_path / "libreoffice-profile-notapid"
    for folder in (dead, alive, unrelated):
        folder.mkdir()
    monkeypatch.setattr(svc, "_process_is_running", lambda pid: pid != 999999)

    assert svc.remove_stale_profiles(tmp_path) == 1
    assert not dead.exists()
    assert alive.exists() and unrelated.exists()
