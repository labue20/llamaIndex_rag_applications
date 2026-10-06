"""
PDF to Word Conversion Module using PyMuPDF

This module provides functions for converting PDF documents 
to Word documents with superior text extraction using PyMuPDF.
"""

import os
import tempfile
import io
from werkzeug.utils import secure_filename
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from config import MAX_UPLOAD_MB, MAX_UPLOAD_BYTES

# Configuration
ALLOWED_EXTENSIONS = {'pdf'}

def allowed_file(filename):
    """Check if file has allowed extension."""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_with_formatting(page):
    """Extract text from PDF page with improved formatting."""
    text_dict = page.get_text("dict")
    page_text = ""
    
    # Process text blocks to maintain formatting
    for block in text_dict["blocks"]:
        if "lines" in block:  # Text block
            block_text = ""
            for line in block["lines"]:
                line_text = ""
                for span in line["spans"]:
                    line_text += span["text"]
                if line_text.strip():
                    block_text += line_text + "\n"
            
            if block_text.strip():
                page_text += block_text + "\n"
    
    return page_text.strip()

def convert_pdf_to_word_document(pdf_content, original_filename):
    """
    Convert PDF content to Word document using PyMuPDF.
    
    Args:
        pdf_content: PDF file content as bytes
        original_filename: Original PDF filename
        
    Returns:
        BytesIO object containing the Word document
    """
    # Open PDF with PyMuPDF
    pdf_document = fitz.open(stream=pdf_content, filetype="pdf")
    
    # Create Word document
    word_doc = Document()
    
    # Add document title based on filename
    title = original_filename.replace('.pdf', '').replace('_', ' ').title()
    title_paragraph = word_doc.add_heading(title, 0)
    title_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # Add document metadata
    word_doc.core_properties.title = title
    word_doc.core_properties.creator = "PDF to Word Converter (PyMuPDF)"
    word_doc.core_properties.subject = f"Converted from {original_filename}"
    
    total_pages = len(pdf_document)
    
    # Process each page
    for page_num in range(total_pages):
        page = pdf_document[page_num]
        
        # Extract text with better formatting
        page_text = extract_text_with_formatting(page)
        
        # Add page header for multi-page documents
        if total_pages > 1:
            page_header = word_doc.add_paragraph(f"--- Page {page_num + 1} ---")
            page_header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            page_header.space_before = Inches(0.2)
            page_header.space_after = Inches(0.1)
        
        # Add page content to Word document
        if page_text:
            # Split into paragraphs and add them
            paragraphs = page_text.split('\n\n')
            for paragraph_text in paragraphs:
                paragraph_text = paragraph_text.strip()
                if paragraph_text:
                    # Split long paragraphs at line breaks
                    lines = paragraph_text.split('\n')
                    for line in lines:
                        line = line.strip()
                        if line:
                            p = word_doc.add_paragraph(line)
                            # Add spacing for better readability
                            if len(lines) == 1:  # Single line paragraphs
                                p.space_after = Inches(0.1)
        else:
            # Add placeholder for pages with no text
            placeholder = word_doc.add_paragraph(f"[Page {page_num + 1} - No text content detected]")
            placeholder.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            placeholder.space_after = Inches(0.2)
        
        # Add page break except for the last page
        if page_num < total_pages - 1:
            word_doc.add_page_break()
    
    # Add footer information
    if total_pages > 1:
        word_doc.add_paragraph()  # Empty line
        footer = word_doc.add_paragraph(
            f"Document converted from {original_filename} using PyMuPDF • "
            f"Total pages: {total_pages} • "
            f"Conversion engine: PyMuPDF + python-docx"
        )
        footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        footer.space_before = Inches(0.3)
    
    # Close PDF document
    pdf_document.close()
    
    # Save Word document to memory
    doc_buffer = io.BytesIO()
    word_doc.save(doc_buffer)
    doc_buffer.seek(0)
    
    return doc_buffer

def validate_pdf_file(uploaded_file):
    """
    Validate uploaded PDF file.
    
    Args:
        uploaded_file: Flask file object
        
    Returns:
        tuple: (is_valid, error_message)
    """
    if not uploaded_file or uploaded_file.filename == '':
        return False, "No file selected"
    
    if not allowed_file(uploaded_file.filename):
        return False, "File must be a PDF"
    
    # Check file size
    uploaded_file.seek(0, 2)  # Seek to end
    file_size = uploaded_file.tell()
    uploaded_file.seek(0)  # Reset to beginning
    
    if file_size > MAX_UPLOAD_BYTES:
        return False, f"File is too large. The maximum size is {MAX_UPLOAD_MB} MB."
    
    if file_size == 0:
        return False, "Empty file provided"
    
    return True, None
